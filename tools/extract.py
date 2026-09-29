# -*- coding: utf-8 -*-
"""抽取參考資料文字內容 -> _extract/*.txt"""
import os, sys, io, glob

BASE = r"D:\全澳中學生資訊科技知識問答比賽\OneDrive_1_2026-9-12"
OUT = r"D:\ICK-Competition\_extract"
os.makedirs(OUT, exist_ok=True)

def dump(name, text):
    p = os.path.join(OUT, name + ".txt")
    with open(p, "w", encoding="utf-8") as f:
        f.write(text)
    print("WROTE", p, len(text), "chars")

def do_pdf(path, name, maxpages=None):
    try:
        import pdfplumber
    except Exception as e:
        print("no pdfplumber", e); return
    out = []
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        if maxpages: n = min(n, maxpages)
        for i in range(n):
            t = pdf.pages[i].extract_text() or ""
            out.append(f"\n===== PAGE {i+1} =====\n" + t)
    dump(name, "\n".join(out))

def do_docx(path, name):
    import docx
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for ti, tb in enumerate(d.tables):
        out.append(f"\n===== TABLE {ti+1} =====")
        for row in tb.rows:
            out.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    dump(name, "\n".join(out))

def do_xlsx(path, name):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n===== SHEET: {ws.title} =====")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells))
    dump(name, "\n".join(out))

tasks = [
    ("pdf", os.path.join(BASE, r"2025年_高中組講義\電腦網絡基礎知識概覽_高中.pdf"), "net"),
    ("pdf", os.path.join(BASE, r"2025年_高中組講義\電腦硬件基礎與資訊科技在日常生活中的應用-高中.pdf"), "hw"),
    ("pdf", os.path.join(BASE, r"2025年_高中組講義\網頁_高中_常用的 HTML, CSS, JS 標籤.pdf"), "web_tags"),
    ("pdf", os.path.join(BASE, r"2025年_高中組講義\網頁_賽前預備課程_高中.pdf"), "web_prep"),
    ("pdf", os.path.join(BASE, r"2024\IN課1.pdf"), "in1"),
    ("pdf", os.path.join(BASE, r"2024\IN課2.pdf"), "in2"),
    ("docx", os.path.join(BASE, r"2024\office基礎.docx"), "office"),
    ("xlsx", os.path.join(BASE, r"2024\Microsoft 365 Apps Online Learning Directory.xlsx"), "m365"),
]

for kind, path, name in tasks:
    if not os.path.exists(path):
        print("MISSING", path); continue
    try:
        print(">>>", kind, name)
        if kind == "pdf": do_pdf(path, name)
        elif kind == "docx": do_docx(path, name)
        elif kind == "xlsx": do_xlsx(path, name)
    except Exception as e:
        print("ERR", name, repr(e))
print("ALL DONE")
